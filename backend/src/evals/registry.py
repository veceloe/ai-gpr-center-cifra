"""Разбор Excel-реестра заказчика в эталонный датасет — задача T017.

В реестре у каждой карточки уже проставлены человеком баллы К1-К6 (или Н1-Н4),
вычисленный индекс, категория и флаг эскалации. Это готовая разметка — другого
эталона у нас нет, и именно на ней измеряется заявленное качество.

Ограничение называем вслух: 46 карточек — не бенчмарк, а верхнеуровневая проверка
на реальных данных заказчика. Расширение эталона — открытый вопрос OQ-08.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

NPA_SHEETS = ("НПА", "Проекты и анонсы")
NEWS_SHEET = "Отраслевые новости"

NPA_CODES = ("К1", "К2", "К3", "К4", "К5", "К6")
NEWS_CODES = ("Н1", "Н2", "Н3", "Н4")


def _find_header_row(rows: list[tuple], marker: str = "№") -> int:
    for index, row in enumerate(rows):
        if row and str(row[0]).strip() == marker:
            return index
    raise ValueError("не найдена строка заголовков (первая ячейка «№»)")


def _column_map(header: tuple, codes: tuple[str, ...]) -> dict[str, int]:
    """Находит колонки критериев по коду в тексте заголовка.

    Заголовки в файле многострочные («К1 Применимость\\n(0–3)»), поэтому ищем вхождение
    кода, а не точное совпадение.
    """
    mapping: dict[str, int] = {}
    for position, cell in enumerate(header):
        text = str(cell or "").strip()
        for code in codes:
            if code in mapping:
                continue
            if text.startswith(code) or f"{code} " in text or text == code:
                mapping[code] = position
    missing = [c for c in codes if c not in mapping]
    if missing:
        raise ValueError(f"в заголовке не найдены критерии: {missing}")
    return mapping


def _find_column(header: tuple, *needles: str) -> int | None:
    for position, cell in enumerate(header):
        text = str(cell or "").lower()
        if all(n.lower() in text for n in needles):
            return position
    return None


def _as_int(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _as_float(value: Any) -> float | None:
    try:
        return round(float(value), 1)
    except (TypeError, ValueError):
        return None


def _parse_sheet(rows: list[tuple], codes: tuple[str, ...], kind: str) -> list[dict]:
    header_index = _find_header_row(rows)
    header = rows[header_index]
    criteria = _column_map(header, codes)

    col_identifier = 1
    col_type = 2
    col_stage = 3
    col_source = 4
    col_index = _find_column(header, "индекс")
    col_category = _find_column(header, "категория")
    col_flag = _find_column(header, "флаг")
    col_final = _find_column(header, "итоговая", "категория")
    col_essence = _find_column(header, "суть")

    parsed: list[dict] = []
    for row in rows[header_index + 1 :]:
        if not row or _as_int(row[0]) is None:
            continue

        scores = {code: _as_int(row[position]) for code, position in criteria.items()}
        if any(value is None for value in scores.values()):
            logger.debug("Карточка %s пропущена: не все баллы проставлены", row[0])
            continue

        flag_raw = str(row[col_flag] or "").strip().lower() if col_flag is not None else ""
        final_category = str(row[col_final] or "").strip() if col_final is not None else ""
        category = str(row[col_category] or "").strip() if col_category is not None else ""

        parsed.append(
            {
                "id": _as_int(row[0]),
                "kind": kind,
                "identifier": str(row[col_identifier] or "").strip(),
                "doc_type": str(row[col_type] or "").strip(),
                "stage": str(row[col_stage] or "").strip(),
                "source_url": str(row[col_source] or "").strip(),
                "text": str(row[col_essence] or "").strip() if col_essence is not None else "",
                "gold_scores": scores,
                "gold_index": _as_float(row[col_index]) if col_index is not None else None,
                "gold_category": category,
                "gold_escalation": flag_raw in {"да", "yes", "true"},
                "gold_final_category": final_category or category,
            }
        )
    return parsed


def parse_registry(path: Path) -> list[dict]:
    """Читает все листы реестра и возвращает плоский список карточек."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    result: list[dict] = []

    for sheet_name in NPA_SHEETS:
        if sheet_name not in workbook.sheetnames:
            logger.warning("Лист %r отсутствует в реестре", sheet_name)
            continue
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        result.extend(_parse_sheet(rows, NPA_CODES, kind="act"))

    if NEWS_SHEET in workbook.sheetnames:
        rows = list(workbook[NEWS_SHEET].iter_rows(values_only=True))
        result.extend(_parse_sheet(rows, NEWS_CODES, kind="news"))

    result.sort(key=lambda r: r["id"])
    return result


def load_dataset(path: Path) -> list[dict]:
    import json

    with path.open(encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]
